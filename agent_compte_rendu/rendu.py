"""Mise en forme du compte rendu en Markdown et en Word (.docx)."""

from pathlib import Path

from .modeles import CompteRendu

NON_PRECISE = "—"


def _entete(cr: CompteRendu) -> list[tuple[str, str]]:
    lignes = [("Date", cr.date or NON_PRECISE)]
    if cr.duree:
        lignes.append(("Durée", cr.duree))
    participants = ", ".join(
        f"{p.nom} ({p.role})" if p.role else p.nom for p in cr.participants
    )
    lignes.append(("Participants", participants or NON_PRECISE))
    lignes.append(("Objet", cr.objet))
    return lignes


def en_markdown(cr: CompteRendu) -> str:
    md = [f"# Compte rendu — {cr.titre}", ""]
    md += [f"**{cle} :** {valeur}  " for cle, valeur in _entete(cr)]
    md += ["", "## Résumé", "", cr.resume, ""]

    md += ["## Points abordés", ""]
    for i, point in enumerate(cr.points_abordes, 1):
        md += [f"### {i}. {point.sujet}", "", point.discussion, ""]
        md += [f"- {c}" for c in point.conclusions]
        md.append("")

    md += ["## Décisions", ""]
    if cr.decisions:
        for d in cr.decisions:
            md.append(f"- **{d.decision}**" + (f" — {d.justification}" if d.justification else ""))
    else:
        md.append("_Aucune décision formelle._")
    md.append("")

    md += ["## Plan d'action", ""]
    if cr.actions:
        md += ["| Action | Responsable | Échéance | Priorité |", "|---|---|---|---|"]
        for a in cr.actions:
            cellules = [a.tache, a.responsable or NON_PRECISE, a.echeance or NON_PRECISE, a.priorite]
            md.append("| " + " | ".join(c.replace("|", "/") for c in cellules) + " |")
    else:
        md.append("_Aucune action identifiée._")
    md.append("")

    if cr.questions_en_suspens:
        md += ["## Questions en suspens", ""]
        md += [f"- {q}" for q in cr.questions_en_suspens]
        md.append("")

    if cr.prochaine_reunion:
        md += ["## Prochaine réunion", "", cr.prochaine_reunion, ""]

    return "\n".join(md)


def en_docx(cr: CompteRendu, chemin: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(f"Compte rendu — {cr.titre}", level=0)
    for cle, valeur in _entete(cr):
        paragraphe = doc.add_paragraph()
        paragraphe.add_run(f"{cle} : ").bold = True
        paragraphe.add_run(valeur)

    doc.add_heading("Résumé", level=1)
    doc.add_paragraph(cr.resume)

    doc.add_heading("Points abordés", level=1)
    for i, point in enumerate(cr.points_abordes, 1):
        doc.add_heading(f"{i}. {point.sujet}", level=2)
        doc.add_paragraph(point.discussion)
        for c in point.conclusions:
            doc.add_paragraph(c, style="List Bullet")

    doc.add_heading("Décisions", level=1)
    if cr.decisions:
        for d in cr.decisions:
            paragraphe = doc.add_paragraph(style="List Bullet")
            paragraphe.add_run(d.decision).bold = True
            if d.justification:
                paragraphe.add_run(f" — {d.justification}")
    else:
        doc.add_paragraph("Aucune décision formelle.")

    doc.add_heading("Plan d'action", level=1)
    if cr.actions:
        tableau = doc.add_table(rows=1, cols=4)
        tableau.style = "Light Grid Accent 1"
        for cellule, titre in zip(tableau.rows[0].cells, ["Action", "Responsable", "Échéance", "Priorité"]):
            cellule.text = titre
        for a in cr.actions:
            ligne = tableau.add_row().cells
            ligne[0].text = a.tache
            ligne[1].text = a.responsable or NON_PRECISE
            ligne[2].text = a.echeance or NON_PRECISE
            ligne[3].text = a.priorite
    else:
        doc.add_paragraph("Aucune action identifiée.")

    if cr.questions_en_suspens:
        doc.add_heading("Questions en suspens", level=1)
        for q in cr.questions_en_suspens:
            doc.add_paragraph(q, style="List Bullet")

    if cr.prochaine_reunion:
        doc.add_heading("Prochaine réunion", level=1)
        doc.add_paragraph(cr.prochaine_reunion)

    doc.save(str(chemin))
